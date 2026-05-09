"""
Document Generator for Map Route Documentation
Generates Word (.docx) and HTML route documents using template-based text.
"""

import os
from datetime import datetime
from typing import Dict, List
import logging
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    print("python-docx not installed. HTML output will be used instead.")

logger = logging.getLogger(__name__)


class RouteDocumentGenerator:
    """Generate Word or HTML documents for route documentation."""

    def __init__(self, use_hf_models: bool = False):
        # HuggingFace models (GPT-2, BART) are intentionally not loaded:
        # they add ~1.6 GB of weight, require GPU/CPU minutes to initialise,
        # and produce lower-quality output for structured navigation text than
        # a simple template.  The template approach is used exclusively.
        if use_hf_models:
            logger.info("HuggingFace model loading is disabled; using template text.")

    # ── Text generation ───────────────────────────────────────────────────────

    def generate_route_description(self, route_data: Dict) -> str:
        return self._generate_basic_description(route_data)

    def _generate_basic_description(self, route_data: Dict) -> str:
        landmark_names = [lm['name'] for lm in route_data.get('landmarks', [])[:3]]
        lm_text = (', '.join(landmark_names) + '.') if landmark_names else 'no specific landmarks detected.'
        direction = route_data.get('direction', 'the indicated direction')
        return (
            f"This route travels from {route_data['source']} to {route_data['destination']} "
            f"in a generally {direction} direction. "
            f"Notable points along the route include {lm_text}"
        )

    def generate_landmark_descriptions(self, landmarks: List[Dict]) -> List[str]:
        type_labels = {
            'start':    'Starting point',
            'end':      'Destination',
            'store':    'Shopping destination',
            'park':     'Recreational area',
            'plaza':    'Commercial complex',
            'hospital': 'Medical facility',
            'school':   'Educational institution',
        }
        descriptions = []
        for i, lm in enumerate(landmarks, 1):
            label = type_labels.get(lm.get('type', ''), 'Point of interest')
            descriptions.append(f"{i}. {lm['name']} — {label}")
        return descriptions

    # ── Word document ─────────────────────────────────────────────────────────

    def create_word_document(self, route_data: Dict, output_path: str) -> str:
        if not DOCX_AVAILABLE:
            logger.warning("python-docx not available; creating HTML instead.")
            return self.create_html_document(route_data, output_path.replace('.docx', '.html'))

        try:
            doc = Document()

            # Title
            title = doc.add_heading('Route Map', level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # Subtitle
            sub = doc.add_paragraph()
            sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = sub.add_run(f'{route_data["source"]}  →  {route_data["destination"]}')
            r.bold = True
            r.font.size = Pt(13)

            # Timestamp
            ts = doc.add_paragraph(f'Generated: {datetime.now().strftime("%B %d, %Y  %I:%M %p")}')
            ts.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph()

            # Route overview
            doc.add_heading('Route Overview', level=1)
            doc.add_paragraph(self.generate_route_description(route_data))

            # Route details table
            doc.add_heading('Route Details', level=1)
            tbl = doc.add_table(rows=2, cols=2)
            tbl.style = 'Light Grid Accent 1'
            tbl.rows[0].cells[0].text = '🟢  Starting Point'
            tbl.rows[0].cells[1].text = route_data['source']
            tbl.rows[1].cells[0].text = '🔴  Destination'
            tbl.rows[1].cells[1].text = route_data['destination']
            for row in tbl.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.bold = True

            doc.add_paragraph()

            # Route visualization
            doc.add_heading('Route Visualization', level=1)
            doc.add_paragraph(
                'The route below is extracted from the uploaded map image. '
                'The green circle marks the starting point (top) and the red '
                'circle marks the destination (bottom).'
            )

            # Embed image if available
            viz = route_data.get('visualization_path', '')
            if viz and os.path.exists(viz):
                try:
                    doc.add_picture(viz, width=Inches(5.5))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception as e:
                    logger.warning(f"Cannot embed image: {e}")
                    doc.add_paragraph(f'[Route image: {viz}]')

            doc.add_paragraph()

            # Landmarks
            if route_data.get('landmarks'):
                doc.add_heading('Landmarks Along the Route', level=1)
                for desc in self.generate_landmark_descriptions(route_data['landmarks']):
                    doc.add_paragraph(desc, style='List Bullet')
                doc.add_paragraph()

            # Disclaimer
            disc_run = doc.add_paragraph().add_run(
                'Disclaimer: This document is generated automatically and is for reference '
                'only. Please verify actual road conditions before travelling.'
            )
            disc_run.font.size = Pt(9)
            disc_run.italic    = True

            doc.save(output_path)
            logger.info(f"Word document saved: {output_path}")
            return output_path

        except Exception as e:
            logger.error(f"Error creating Word document: {e}")
            return self.create_html_document(route_data, output_path.replace('.docx', '.html'))

    # ── HTML document ─────────────────────────────────────────────────────────

    def create_html_document(self, route_data: Dict, output_path: str) -> str:
        route_desc = self.generate_route_description(route_data)
        landmark_items = ''.join(
            f'<li>{d}</li>'
            for d in self.generate_landmark_descriptions(route_data.get('landmarks', []))
        )

        html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>Route Map — {route_data['source']} to {route_data['destination']}</title>
  <style>
    body      {{ font-family: Arial, sans-serif; max-width: 800px; margin: 0 auto; padding: 20px; line-height: 1.6; }}
    h1        {{ text-align: center; }}
    .subtitle {{ text-align: center; font-size: 1.2rem; font-weight: bold; margin-bottom: .25rem; }}
    .date     {{ text-align: right; color: #666; margin-bottom: 1.5rem; }}
    .info-box {{ background: #f0f8ff; padding: 15px; border-radius: 5px; margin: 1rem 0; }}
    .landmark {{ background: #f9f9f9; padding: 15px; border-radius: 5px; }}
    .disc     {{ background: #fff3cd; padding: 15px; border-radius: 5px; margin-top: 2rem; border-left: 4px solid #ffa500; font-size: .9rem; }}
    table     {{ width: 100%; border-collapse: collapse; margin: 1rem 0; }}
    td        {{ padding: 8px 12px; border: 1px solid #ddd; }}
    td:first-child {{ font-weight: bold; width: 40%; background: #f5f5f5; }}
    img       {{ max-width: 100%; height: auto; display: block; margin: 1rem auto; border: 1px solid #ddd; border-radius: 5px; }}
  </style>
</head>
<body>
  <h1>Route Map</h1>
  <p class="subtitle">{route_data['source']} → {route_data['destination']}</p>
  <p class="date">Generated: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}</p>

  <h2>Route Overview</h2>
  <p>{route_desc}</p>

  <h2>Route Details</h2>
  <table>
    <tr><td>🟢 Starting Point</td><td>{route_data['source']}</td></tr>
    <tr><td>🔴 Destination</td>   <td>{route_data['destination']}</td></tr>
  </table>

  {('<h2>Route Visualization</h2>'
    '<p>🟢 Green circle — Starting point (top)<br>'
    '🔵 Blue line — Route path<br>'
    '🔴 Red circle — Destination (bottom)</p>'
    f'<img src="{route_data.get("visualization_path","")}" alt="Route visualization">')
   if route_data.get("visualization_path") else ''}

  {('<h2>Landmarks</h2><ul class="landmark">' + landmark_items + '</ul>')
   if landmark_items else ''}

  <div class="disc">
    <strong>Disclaimer:</strong> This document is generated automatically and is for reference only.
    Please verify actual road conditions before travelling.
  </div>
</body>
</html>"""

        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(html)
        logger.info(f"HTML document saved: {output_path}")
        return output_path

    # ── Public entry point ────────────────────────────────────────────────────

    def generate_route_document(self, route_data: Dict,
                                output_dir: str, format: str = 'docx') -> str:
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        src = route_data['source'].replace(' ', '_')
        dst = route_data['destination'].replace(' ', '_')

        if format.lower() == 'docx':
            path = os.path.join(output_dir, f'route_{src}_{dst}_{timestamp}.docx')
            return self.create_word_document(route_data, path)

        if format.lower() == 'html':
            path = os.path.join(output_dir, f'route_{src}_{dst}_{timestamp}.html')
            return self.create_html_document(route_data, path)

        raise ValueError(f"Unsupported format: {format}")


def main():
    route_data = {
        'source':      'Home',
        'destination': 'NTR Garden',
        'direction':   'North-East',
        'landmarks':   [
            {'name': 'Shopping Mall', 'type': 'store'},
            {'name': 'City Park',     'type': 'park'},
        ],
        'visualization_path': 'output/route_visualization.png',
    }

    generator = RouteDocumentGenerator()
    output_dir = 'output'
    Path(output_dir).mkdir(parents=True, exist_ok=True)

    try:
        path = generator.generate_route_document(route_data, output_dir, 'docx')
        print(f"Word document: {path}")
    except Exception as e:
        print(f"Error: {e}")
        path = generator.generate_route_document(route_data, output_dir, 'html')
        print(f"HTML document: {path}")


if __name__ == "__main__":
    main()
